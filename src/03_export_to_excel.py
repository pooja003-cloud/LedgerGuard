"""
LedgerGuard: Build the analysis workbook
========================================

Creates  excel/LedgerGuard_Analysis_generated.xlsx  from data/ledgerguard.db.
Every figure is a live Excel formula over the exported data, so the workbook
recalculates if the data changes.

    About           what the workbook contains
    Summary         KPI figures, exposure by type and by retailer
    Gross_to_Net    revenue bridge from contract value to net revenue, with tie-out
    Returns_Reserve faulty-batch analysis and holiday downside scenario
    Exceptions      exception register (from the v_exceptions view)
    Shipments       all invoices, with month and expected return rate
    Contract_Terms  agreed price and limits per retailer and product

The presented version (dashboard layout and charts, built by hand in Google
Sheets) is excel/LedgerGuard_Analysis_completed.xlsx. This script reproduces
all of its numbers.

Run:  python src/03_export_to_excel.py
"""

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DB_PATH = PROJECT_ROOT / "data" / "ledgerguard.db"
VIEWS_SQL = PROJECT_ROOT / "sql" / "02_exceptions.sql"
OUT_PATH = PROJECT_ROOT / "excel" / "LedgerGuard_Analysis_generated.xlsx"

FONT = "Arial"
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
NUM = '#,##0;(#,##0);"-"'
NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SUBTOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="808080")
INPUT_FONT = Font(name=FONT, size=10, color="0000FF")   # blue = typed-in assumption


def load_data():
    """Read tables and the exception view from an in-memory copy of the database."""
    src = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)
    conn = sqlite3.connect(":memory:")
    src.backup(conn)
    src.close()
    if not conn.execute("SELECT 1 FROM sqlite_master WHERE name='v_exceptions'").fetchone():
        conn.executescript(VIEWS_SQL.read_text())
    shipments = pd.read_sql("""
        SELECT s.*, c.expected_return_rate
        FROM shipments s JOIN contract_terms c USING (retailer, product)
        ORDER BY s.invoice_date, s.invoice_id""", conn)
    terms = pd.read_sql("SELECT * FROM contract_terms", conn)
    exceptions = pd.read_sql("SELECT * FROM v_exceptions ORDER BY impact_usd DESC", conn)
    conn.close()
    return shipments, terms, exceptions


def font(cell, **kw):
    cell.font = Font(name=FONT, size=kw.pop("size", 10), **kw)


def header_row(ws, row, labels):
    for i, h in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=h)
        font(c, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL


def title(ws, text, subtitle):
    ws["A1"] = text
    font(ws["A1"], size=14, bold=True, color=NAVY)
    ws["A2"] = subtitle
    font(ws["A2"], italic=True, color="595959")


def write_table(ws, df, table_name, formats):
    """Write a DataFrame as a formatted Excel Table. Returns the last data row."""
    df = df.copy()
    df.columns = [c.title().replace("Usd", "USD").replace("Id", "ID") for c in df.columns]
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    for cell in ws[1]:
        font(cell, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            font(cell)
    for col_name, fmt in formats.items():
        idx = list(df.columns).index(col_name) + 1
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            cell.number_format = fmt
    for i, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
    table = Table(displayName=table_name,
                  ref=f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    return len(df) + 1


def build_about(ws):
    title(ws, "LedgerGuard: Revenue Leakage Analysis",
          "Simulated data for Northwind Devices (fictional), a PC hardware vendor selling to Best Buy, Amazon and Costco "
          "(Sep 2025 - Oct 15, 2026). Prices and contract terms are illustrative.")
    lines = [
        ("Sheets", True),
        ("Summary  -  headline figures, exposure by type and by retailer", False),
        ("Gross_to_Net  -  revenue bridge from contract value to net revenue, with tie-out check", False),
        ("Returns_Reserve  -  faulty-batch analysis and holiday downside scenario", False),
        ("Exceptions  -  every invoice that broke a contract rule, with $ impact and action", False),
        ("Shipments  -  all 1,000 invoices (source data)", False),
        ("Contract_Terms  -  agreed price and limits per retailer and product", False),
        ("", False),
        ("Conventions", True),
        ("Blue text = typed-in assumption. Black text = formula. Negative numbers in parentheses.", False),
        ("", False),
        ("Generated by src/03_export_to_excel.py from data/ledgerguard.db", False),
    ]
    for i, (text, bold) in enumerate(lines, start=4):
        font(ws.cell(row=i, column=1, value=text), size=11, bold=bold)
    ws.column_dimensions["A"].width = 110


def build_gross_to_net(ws, r):
    S = lambda col: f"Shipments!${col}$2:${col}${r}"   # noqa: E731
    title(ws, "Gross-to-Net Revenue Bridge",
          "Sep 2025 - Oct 15, 2026  |  All retailers  |  USD  |  Deductions shown as negatives")
    header_row(ws, 4, ["Line item", "Amount ($)", "% of contract value"])
    lines = [
        ("Gross invoiced at contract price", f"=SUMPRODUCT({S('E')},{S('F')})", True),
        ("Less: pricing errors (billed below contract)", f"=SUM({S('H')})-B5", False),
        ("Gross invoiced (as billed)", "=B5+B6", True),
        ("Less: markdown allowances", f"=-SUM({S('I')})", False),
        ("Less: chargebacks - routine", f"=-(SUM({S('J')})+B10)", False),
        ("Less: chargebacks - disputed",
         '=-SUMIFS(Exceptions!$F:$F,Exceptions!$E:$E,"CHARGEBACK_OVER_CAP")', False),
        ("Less: returns credits", f"=-SUMPRODUCT({S('L')},{S('G')})", False),
        ("Net revenue", "=SUM(B7:B11)", True),
    ]
    for i, (label, formula, subtotal) in enumerate(lines, start=5):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=formula).number_format = USD
        ws.cell(row=i, column=3, value=f"=B{i}/$B$5").number_format = PCT
        for col in range(1, 4):
            cell = ws.cell(row=i, column=col)
            font(cell, bold=subtotal)
            cell.border = Border(bottom=THIN)
            if subtotal:
                cell.fill = SUBTOTAL_FILL
    ws["A14"] = "Check: net revenue calculated directly from Shipments"
    ws["B14"] = f"=SUM({S('H')})-SUM({S('I')})-SUM({S('J')})-SUMPRODUCT({S('L')},{S('G')})"
    ws["A15"] = "Difference (should be $0)"
    ws["B15"] = "=B12-B14"
    ws["A17"] = ("Returns credits include the faulty March 2026 AI laptop batch "
                 "(see Returns_Reserve); it is already reflected in net revenue.")
    for ref in ("A14", "A15", "A17"):
        font(ws[ref], italic=True, color="595959")
    for ref in ("B14", "B15"):
        font(ws[ref])
        ws[ref].number_format = USD
    for col, w in zip("ABC", (50, 18, 20)):
        ws.column_dimensions[col].width = w
    add_waterfall(ws)


def add_waterfall(ws):
    """Waterfall chart as a stacked column chart: an invisible 'base' series lifts
    each deduction to where the running total stands. All helper cells are formulas."""
    ws["E4"], ws["F4"], ws["G4"], ws["H4"] = "Chart data", "Base", "Total", "Deduction"
    for ref in ("E4", "F4", "G4", "H4"):
        font(ws[ref], bold=True, color="595959")
    rows = [  # label, base, total, deduction (positive heights)
        ("Gross at contract", "=0", "=B5", "=0"),
        ("Pricing errors", "=B7", "=0", "=-B6"),
        ("Gross as billed", "=0", "=B7", "=0"),
        ("Markdowns", "=B7+B8", "=0", "=-B8"),
        ("Routine chargebacks", "=B7+B8+B9", "=0", "=-B9"),
        ("Disputed chargebacks", "=B7+B8+B9+B10", "=0", "=-B10"),
        ("Returns credits", "=B12", "=0", "=-B11"),
        ("Net revenue", "=0", "=B12", "=0"),
    ]
    for i, vals in enumerate(rows, start=5):
        for j, v in enumerate(vals):
            c = ws.cell(row=i, column=5 + j, value=v)
            font(c, color="595959")
            if j:
                c.number_format = USD

    chart = BarChart()
    chart.type, chart.grouping, chart.overlap, chart.gapWidth = "col", "stacked", 100, 60
    chart.title = "Gross-to-Net Revenue Bridge (axis from $200M)"
    chart.add_data(Reference(ws, min_col=6, max_col=8, min_row=4, max_row=12), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=5, min_row=5, max_row=12))
    base, total, deduct = chart.series
    base.graphicalProperties = GraphicalProperties(noFill=True)
    base.graphicalProperties.line.noFill = True
    total.graphicalProperties = GraphicalProperties(solidFill="3D6CB3")
    deduct.graphicalProperties = GraphicalProperties(solidFill="B23A3A")
    chart.y_axis.scaling.min = 200_000_000
    chart.y_axis.scaling.max = 230_000_000
    chart.y_axis.majorUnit = 5_000_000
    chart.y_axis.numFmt = '$#,##0,,"M"'
    chart.y_axis.delete = False
    chart.x_axis.delete = False
    chart.legend = None
    chart.height, chart.width = 9, 22
    ws.add_chart(chart, "A19")


def build_returns_reserve(ws, r):
    S = lambda col: f"Shipments!${col}$2:${col}${r}"   # noqa: E731
    match = f"(Shipments!$D$2:$D${r}=$B$6)*(Shipments!$M$2:$M${r}=$B$5)"
    title(ws, "Returns Reserve Analysis: AI Laptop 14",
          "Customer returns are a refund liability (returns reserve under ASC 606), separate from warranty.")

    def row(n, label, value, fmt, bold=False, section=False):
        a = ws.cell(row=n, column=1, value=label)
        if section:
            font(a, bold=True, color="FFFFFF")
            a.fill = HEADER_FILL
            ws.cell(row=n, column=2).fill = HEADER_FILL
            return
        font(a, bold=bold)
        if value is not None:
            b = ws.cell(row=n, column=2, value=value)
            b.number_format = fmt
            font(b, bold=bold)

    row(4, "Inputs", None, None, section=True)
    row(5, "Faulty batch month", "2026-03", "@")
    row(6, "Product", "AI_Laptop_14", "@")
    row(7, "Return-rate trigger (x expected)", 1.5, '0.0"x"')
    for ref in ("B5", "B6", "B7"):
        ws[ref].font = INPUT_FONT

    row(9, "Faulty batch: actual", None, None, section=True)
    row(10, "Units shipped", f"=SUMIFS({S('E')},{S('D')},B6,{S('M')},B5)", NUM)
    row(11, "Units returned", f"=SUMIFS({S('L')},{S('D')},B6,{S('M')},B5)", NUM)
    row(12, "Actual return rate", "=B11/B10", PCT2)
    row(13, "Expected returns (contract rates, units-weighted)",
        f"=SUMPRODUCT({match}*{S('E')}*{S('N')})", NUM)
    row(14, "Expected return rate", "=B13/B10", PCT2)
    row(15, "Actual vs expected (x)", "=B12/B14", '0.0"x"')
    row(16, "Excess returns (units)", "=B11-B13", NUM)
    row(17, "Average contract price", f"=SUMPRODUCT({match}*{S('E')}*{S('F')})/B10", USD2)
    row(18, "Excess returns at contract price (upper bound)", "=B16*B17", USD, bold=True)

    row(20, "Holiday 2026 downside scenario", None, None, section=True)
    row(21, "Forecast units Nov-Dec 2026 (= Nov-Dec 2025 actual)",
        f'=SUMIFS({S("E")},{S("D")},B6,{S("M")},"2025-11")+SUMIFS({S("E")},{S("D")},B6,{S("M")},"2025-12")', NUM)
    row(22, "Extra returns if the March rate recurred, at contract price",
        "=B21*(B12-B14)*B17", USD, bold=True)

    notes = [
        "Notes",
        "1. The excess returns are already included in 'Less: returns credits' on Gross_to_Net. "
        "They reduce revenue; no additional charge is needed.",
        "2. Returns were back to plan in April 2026, so this is treated as a one-off variance with "
        "no change to the ongoing returns reserve rate.",
        "3. Values at full contract price are an upper bound: returned units keep some recovery value "
        "(refurbish and resell, or return to the supplier).",
    ]
    for i, text in enumerate(notes, start=24):
        font(ws.cell(row=i, column=1, value=text), bold=(i == 24), italic=(i > 24), color="595959")
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 18


def build_summary(ws, last_exc):
    E = lambda col: f"Exceptions!${col}$2:${col}${last_exc}"   # noqa: E731
    title(ws, "Revenue Leakage Summary",
          "Sep 2025 - Oct 15, 2026  |  USD  |  Simulated data")
    kpis = [
        ("Total exposure", f"=SUM({E('F')})", USD),
        ("Recoverable (up to)", f'=SUMIFS({E("F")},{E("G")},"Yes")', USD),
        ("Not recoverable", f'=SUMIFS({E("F")},{E("G")},"No")', USD),
        ("% of gross billed", "=B5/Gross_to_Net!B7", PCT2),
        ("Holiday downside if defect recurred", "=Returns_Reserve!B22", USD),
    ]
    header_row(ws, 4, ["Headline figure", "Amount"])
    for i, (label, formula, fmt) in enumerate(kpis, start=5):
        font(ws.cell(row=i, column=1, value=label), bold=True)
        c = ws.cell(row=i, column=2, value=formula)
        c.number_format = fmt
        font(c, bold=True)

    header_row(ws, 11, ["Exposure by type", "Invoices", "Amount", "Recoverable?"])
    types = [("PRICE_BELOW_CONTRACT", "Yes"), ("CHARGEBACK_OVER_CAP", "If dispute succeeds"),
             ("RETURN_SPIKE", "No"), ("MARKDOWN_OVER_CAP", "Yes")]
    for i, (t, rec) in enumerate(types, start=12):
        font(ws.cell(row=i, column=1, value=t))
        c = ws.cell(row=i, column=2, value=f'=COUNTIFS({E("E")},A{i})')
        c.number_format = NUM
        c = ws.cell(row=i, column=3, value=f'=SUMIFS({E("F")},{E("E")},A{i})')
        c.number_format = USD
        font(ws.cell(row=i, column=4, value=rec))
    font(ws.cell(row=16, column=1, value="Total"), bold=True)
    for col, fmt in (("B", NUM), ("C", USD)):
        c = ws[f"{col}16"]
        c.value = f"=SUM({col}12:{col}15)"
        c.number_format = fmt
        font(c, bold=True)

    header_row(ws, 18, ["Exposure by retailer", "Pricing", "Chargebacks", "Returns", "Total"])
    cols = {"B": "PRICE_BELOW_CONTRACT", "C": "CHARGEBACK_OVER_CAP", "D": "RETURN_SPIKE"}
    for i, ret in enumerate(["Amazon_US", "BestBuy_US", "Costco_Wholesale"], start=19):
        font(ws.cell(row=i, column=1, value=ret))
        for col, t in cols.items():
            c = ws[f"{col}{i}"]
            c.value = f'=SUMIFS({E("F")},{E("C")},$A{i},{E("E")},"{t}")'
            c.number_format = USD
        ws[f"E{i}"] = f"=SUM(B{i}:D{i})"
        ws[f"E{i}"].number_format = USD
    for ref in ("B19", "B20", "B21", "C19", "C20", "C21", "D19", "D20", "D21", "E19", "E20", "E21"):
        font(ws[ref])
    font(ws.cell(row=22, column=1, value="Total"), bold=True)
    for col in "BCDE":
        c = ws[f"{col}22"]
        c.value = f"=SUM({col}19:{col}21)"
        c.number_format = USD
        font(c, bold=True)

    ws["A24"] = "Check: by-retailer total equals total exposure (should be $0)"
    ws["B24"] = "=E22-B5"
    font(ws["A24"], italic=True, color="595959")
    ws["B24"].number_format = USD
    font(ws["B24"])
    for col, w in zip("ABCDE", (44, 18, 16, 16, 16)):
        ws.column_dimensions[col].width = w


def main():
    shipments, terms, exceptions = load_data()
    shipments["invoice_date"] = pd.to_datetime(shipments["invoice_date"])
    shipments.insert(12, "month", shipments["invoice_date"].dt.strftime("%Y-%m"))
    exceptions["invoice_date"] = pd.to_datetime(exceptions["invoice_date"])
    exceptions.insert(6, "recoverable",
                      exceptions["exception_type"].map(lambda t: "No" if t == "RETURN_SPIKE" else "Yes"))

    wb = Workbook()
    wb.active.title = "About"
    build_about(wb.active)
    ws_sum = wb.create_sheet("Summary")
    ws_g2n = wb.create_sheet("Gross_to_Net")
    ws_res = wb.create_sheet("Returns_Reserve")
    ws_exc = wb.create_sheet("Exceptions")
    ws_ship = wb.create_sheet("Shipments")
    ws_terms = wb.create_sheet("Contract_Terms")

    last_exc = write_table(ws_exc, exceptions, "tblExceptions",
                           {"Invoice_Date": "yyyy-mm-dd", "Impact_USD": USD2})
    last_ship = write_table(ws_ship, shipments, "tblShipments", {
        "Invoice_Date": "yyyy-mm-dd", "Contract_Unit_Price": USD2, "Billed_Unit_Price": USD2,
        "Gross_Invoice": USD2, "Markdown_Allowance": USD2, "Chargeback": USD2,
        "Units_Shipped": "#,##0", "Units_Returned": "#,##0", "Expected_Return_Rate": "0.0%"})
    write_table(ws_terms, terms, "tblContractTerms", {
        "Contract_Unit_Price": USD2, "Max_Markdown_Pct": "0.0%",
        "Max_Chargeback_Pct": "0.0%", "Expected_Return_Rate": "0.0%"})

    build_gross_to_net(ws_g2n, last_ship)
    build_returns_reserve(ws_res, last_ship)
    build_summary(ws_sum, last_exc)
    wb.active = 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Saved {OUT_PATH.relative_to(PROJECT_ROOT)}: "
          f"{len(shipments)} shipments, {len(exceptions)} exceptions")


if __name__ == "__main__":
    main()
